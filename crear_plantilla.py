"""
Genera la plantilla Word por defecto para el informe de inspección estructural.

La plantilla usa etiquetas Jinja (docxtpl). Puede abrirse y modificarse en Word
libremente (logos, estilos, textos...) siempre que se respeten las etiquetas
{{ ... }}, {%p ... %} y {%tr ... %}.

Uso:
    python crear_plantilla.py [ruta_salida.docx]
"""
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

SALIDA_POR_DEFECTO = Path(__file__).parent / "plantillas" / "plantilla_inspeccion.docx"


def _parrafo(doc, texto, centrado=False, negrita=False, tam=None):
    p = doc.add_paragraph()
    run = p.add_run(texto)
    run.bold = negrita
    if tam:
        run.font.size = Pt(tam)
    if centrado:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def crear_plantilla(salida: Path = SALIDA_POR_DEFECTO) -> Path:
    doc = Document()

    # ---------------- Portada ----------------
    _parrafo(doc, "INFORME DE INSPECCIÓN ESTRUCTURAL", centrado=True, negrita=True, tam=20)
    _parrafo(doc, "{{ obra }}", centrado=True, negrita=True, tam=14)
    # Imagen de portada: se rellena con un archivo llamado "portada.jpg" (o .png)
    _parrafo(doc, "{{ img_portada }}", centrado=True)

    doc.add_heading("1. Datos generales", level=1)
    datos = [
        ("Obra / Estructura", "{{ obra }}"),
        ("Ubicación", "{{ ubicacion }}"),
        ("Cliente", "{{ cliente }}"),
        ("Inspector", "{{ inspector }}"),
        ("Fecha de inspección", "{{ fecha }}"),
        ("N.º de defectos registrados", "{{ total_defectos }}"),
    ]
    tabla = doc.add_table(rows=0, cols=2)
    tabla.style = "Table Grid"
    for etiqueta, valor in datos:
        fila = tabla.add_row().cells
        fila[0].text = etiqueta
        fila[0].paragraphs[0].runs[0].bold = True
        fila[1].text = valor

    doc.add_heading("2. Objeto y alcance", level=1)
    doc.add_paragraph(
        "El presente informe recoge los resultados de la inspección visual realizada "
        "sobre la estructura indicada, documentando fotográficamente los defectos "
        "detectados, su localización y su nivel de severidad."
    )
    # Plano o croquis de situación: archivo "ubicacion.jpg" (opcional)
    _parrafo(doc, "{{ img_ubicacion }}", centrado=True)

    # ---------------- Registro de defectos ----------------
    doc.add_heading("3. Registro fotográfico de defectos", level=1)
    doc.add_paragraph("{%p for d in defectos %}")
    _parrafo(doc, "Defecto N.º {{ d.num }} – {{ d.tipo }}", negrita=True, tam=12)

    ficha = doc.add_table(rows=0, cols=2)
    ficha.style = "Table Grid"
    for etiqueta, valor in [
        ("Tipo de defecto", "{{ d.tipo }}"),
        ("Elemento", "{{ d.elemento }}"),
        ("Ubicación", "{{ d.ubicacion }}"),
        ("Severidad", "{{ d.severidad }}"),
        ("Observaciones", "{{ d.observaciones }}"),
        ("Archivo", "{{ d.archivo }}"),
    ]:
        fila = ficha.add_row().cells
        fila[0].text = etiqueta
        fila[0].paragraphs[0].runs[0].bold = True
        fila[1].text = valor

    _parrafo(doc, "{{ d.imagen }}", centrado=True)
    _parrafo(doc, "Fotografía {{ d.num }}. {{ d.tipo }} – {{ d.elemento }}", centrado=True, tam=9)
    doc.add_paragraph("{%p endfor %}")

    # ---------------- Resumen ----------------
    doc.add_heading("4. Resumen de defectos", level=1)
    doc.add_paragraph("Número total de defectos hallados por tipo:")

    resumen = doc.add_table(rows=1, cols=3)
    resumen.style = "Table Grid"
    resumen.alignment = WD_TABLE_ALIGNMENT.CENTER
    for celda, texto in zip(resumen.rows[0].cells, ["Tipo de defecto", "Cantidad", "%"]):
        celda.text = texto
        celda.paragraphs[0].runs[0].bold = True
    # Filas generadas por bucle ({%tr %} debe ir solo en su propia fila)
    resumen.add_row().cells[0].text = "{%tr for r in resumen %}"
    fila = resumen.add_row().cells
    fila[0].text, fila[1].text, fila[2].text = "{{ r.tipo }}", "{{ r.cantidad }}", "{{ r.porcentaje }}"
    resumen.add_row().cells[0].text = "{%tr endfor %}"
    fila = resumen.add_row().cells
    fila[0].text, fila[1].text, fila[2].text = "TOTAL", "{{ total_defectos }}", "100 %"
    for celda in fila:
        celda.paragraphs[0].runs[0].bold = True

    doc.add_paragraph("")
    doc.add_paragraph("Distribución por severidad:")
    sev = doc.add_table(rows=1, cols=2)
    sev.style = "Table Grid"
    sev.alignment = WD_TABLE_ALIGNMENT.CENTER
    for celda, texto in zip(sev.rows[0].cells, ["Severidad", "Cantidad"]):
        celda.text = texto
        celda.paragraphs[0].runs[0].bold = True
    sev.add_row().cells[0].text = "{%tr for s in resumen_severidad %}"
    fila = sev.add_row().cells
    fila[0].text, fila[1].text = "{{ s.severidad }}", "{{ s.cantidad }}"
    sev.add_row().cells[0].text = "{%tr endfor %}"

    doc.add_heading("5. Conclusiones", level=1)
    doc.add_paragraph("{{ conclusiones }}")

    doc.add_paragraph("")
    _parrafo(doc, "Fdo.: {{ inspector }}")
    _parrafo(doc, "Fecha: {{ fecha }}")

    salida = Path(salida)
    salida.parent.mkdir(parents=True, exist_ok=True)
    doc.save(salida)
    return salida


if __name__ == "__main__":
    destino = Path(sys.argv[1]) if len(sys.argv) > 1 else SALIDA_POR_DEFECTO
    print(f"Plantilla creada en: {crear_plantilla(destino)}")

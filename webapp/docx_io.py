"""Lectura de plantillas (docx/txt/md/pdf) y generación del informe Word a partir de la spec."""
import io
from pathlib import Path

from docx import Document
from docx.shared import Mm
from docx.table import Table
from docx.text.paragraph import Paragraph
from PIL import Image, ImageOps

MAX_PLANTILLA_CHARS = 30000
LADO_DOCX_PX = 1600
LADO_LLM_PX = 1024
EXT_IMAGEN = {".jpg", ".jpeg", ".png", ".bmp", ".tif", ".tiff", ".gif", ".webp"}


def jpeg_bytes(origen, lado_max: int, calidad: int = 85) -> bytes:
    """Imagen (ruta o bytes) -> JPEG con orientación EXIF corregida y lado máximo."""
    img = Image.open(io.BytesIO(origen) if isinstance(origen, bytes) else origen)
    img = ImageOps.exif_transpose(img)
    if img.mode != "RGB":
        img = img.convert("RGB")
    img.thumbnail((lado_max, lado_max))
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=calidad)
    return buf.getvalue()


def _docx_en_blanco() -> bytes:
    buf = io.BytesIO()
    Document().save(buf)
    return buf.getvalue()


def _pdf_a_texto(datos: bytes) -> str:
    from pypdf import PdfReader
    return "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(datos)).pages)


def cargar_plantilla(nombre: str, datos: bytes) -> tuple[bytes, str]:
    """Devuelve (docx_base, texto_estructura). Para docx el base conserva estilos, cabeceras y pies."""
    ext = Path(nombre).suffix.lower()
    if ext == ".docx":
        Document(io.BytesIO(datos))  # valida
        return datos, estructura_docx(datos)
    if ext in {".txt", ".md"}:
        texto = datos.decode("utf-8", errors="replace")
    elif ext == ".pdf":
        texto = _pdf_a_texto(datos)
    else:
        raise ValueError("Formato de plantilla no admitido (use .docx, .pdf, .txt o .md)")
    return _docx_en_blanco(), texto[:MAX_PLANTILLA_CHARS]


def _iter_cuerpo(doc):
    for el in doc.element.body.iterchildren():
        if el.tag.endswith("}p"):
            yield Paragraph(el, doc)
        elif el.tag.endswith("}tbl"):
            yield Table(el, doc)


def estructura_docx(datos: bytes) -> str:
    """Descripción textual del documento: estilo y texto de cada párrafo, tablas, cabeceras y pies."""
    doc = Document(io.BytesIO(datos))
    lineas = []
    for sec in doc.sections[:1]:
        cab = " / ".join(p.text for p in sec.header.paragraphs if p.text.strip())
        pie = " / ".join(p.text for p in sec.footer.paragraphs if p.text.strip())
        if cab:
            lineas.append(f"[CABECERA] {cab}")
        if pie:
            lineas.append(f"[PIE] {pie}")
    for el in _iter_cuerpo(doc):
        if isinstance(el, Paragraph):
            if el.text.strip():
                lineas.append(f"[{el.style.name}] {el.text.strip()}")
        else:
            lineas.append("[TABLA]")
            for fila in el.rows:
                lineas.append("  | " + " | ".join(c.text.strip().replace("\n", " ") for c in fila.cells))
    return "\n".join(lineas)[:MAX_PLANTILLA_CHARS]


def _estilo(doc, *nombres):
    for n in nombres:
        try:
            return doc.styles[n]
        except KeyError:
            continue
    return None


def _ancho_util_mm(doc) -> float:
    s = doc.sections[0]
    return max(50.0, (s.page_width - s.left_margin - s.right_margin) / 36000)


def renderizar(base: bytes, spec: dict, ruta_foto) -> bytes:
    """Construye el informe sobre la plantilla base (se vacía el cuerpo, se conservan estilos/cabeceras)."""
    doc = Document(io.BytesIO(base))
    cuerpo = doc.element.body
    for el in list(cuerpo.iterchildren()):
        if not el.tag.endswith("}sectPr"):
            cuerpo.remove(el)

    ancho = min(150.0, _ancho_util_mm(doc))
    st_titulo = _estilo(doc, "Title")
    if st_titulo:
        doc.add_paragraph(spec["title"], style=st_titulo)
    else:
        doc.add_paragraph().add_run(spec["title"]).bold = True

    for b in spec["blocks"]:
        t = b["type"]
        if t == "heading":
            st = _estilo(doc, f"Heading {b['level']}")
            if st:
                doc.add_paragraph(b["text"], style=st)
            else:
                doc.add_paragraph().add_run(b["text"]).bold = True
        elif t == "paragraph":
            doc.add_paragraph(b["text"])
        elif t == "bullets":
            st = _estilo(doc, "List Bullet")
            for i in b["items"]:
                doc.add_paragraph(i if st else f"• {i}", style=st)
        elif t == "table":
            n = len(b["header"]) or len(b["rows"][0])
            tabla = doc.add_table(rows=0, cols=n)
            st = _estilo(doc, "Table Grid")
            if st:
                tabla.style = st
            if b["header"]:
                for c, txt in zip(tabla.add_row().cells, b["header"]):
                    c.paragraphs[0].add_run(txt).bold = True
            for fila in b["rows"]:
                for c, txt in zip(tabla.add_row().cells, fila):
                    c.text = txt
            doc.add_paragraph()
        elif t == "photo":
            p = doc.add_paragraph()
            p.alignment = 1
            p.add_run().add_picture(io.BytesIO(jpeg_bytes(ruta_foto(b["file"]), LADO_DOCX_PX)), width=Mm(ancho))
            if b["caption"]:
                st = _estilo(doc, "Caption")
                cap = doc.add_paragraph(style=st)
                cap.alignment = 1
                r = cap.add_run(b["caption"])
                r.italic = True
        elif t == "page_break":
            doc.add_page_break()

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

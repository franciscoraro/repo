import io
import time

import pytest
from docx import Document
from fastapi.testclient import TestClient
from PIL import Image


@pytest.fixture()
def cliente(tmp_path, monkeypatch):
    monkeypatch.setenv("INFORMES_DATA_DIR", str(tmp_path))
    from webapp import llm, store
    monkeypatch.setattr(store, "DATA", tmp_path)
    llamadas = []

    def falso(**kw):
        llamadas.append(kw)
        ids = [f for f, _ in kw["fotos"]]
        bloques = [{"type": "heading", "level": 1, "text": "Resumen"},
                   {"type": "paragraph", "text": "feedback: %s" % kw["feedback"]},
                   {"type": "table", "header": ["A", "B"], "rows": [["1"]]},
                   {"type": "photo", "file": ids[0], "caption": "primera"},
                   {"type": "photo", "file": "no_existe.jpg", "caption": "x"}]
        return {"title": "Informe test", "blocks": bloques}

    monkeypatch.setattr(llm, "generar_spec", falso)
    from webapp.main import app
    c = TestClient(app)
    c.llamadas = llamadas
    return c


def _png():
    b = io.BytesIO()
    Image.new("RGB", (40, 30), "red").save(b, "PNG")
    return b.getvalue()


def _plantilla():
    d = Document()
    d.add_heading("Informe de inspección", 1)
    d.add_paragraph("Texto fijo")
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def esperar(c, jid):
    for _ in range(100):
        j = c.get(f"/api/trabajos/{jid}").json()
        if j["estado"] not in ("generando", "preparando"):
            return j
        time.sleep(0.05)
    raise AssertionError("timeout")


def crear(c, **extra):
    datos = {"nombre": "Puente", "instrucciones": "Describe defectos",
             "rutas": ["Fisuras/a.png", "b.png", "nota.txt"], **extra}
    files = [("fotos", ("a.png", _png(), "image/png")), ("fotos", ("b.png", _png(), "image/png")),
             ("fotos", ("nota.txt", b"x", "text/plain"))]
    if "plantilla_id" not in extra:
        files.append(("plantilla_archivo", ("p.docx", _plantilla(), "application/octet-stream")))
    return c.post("/api/trabajos", data=datos, files=files)


def test_flujo_completo(cliente):
    r = crear(cliente)
    assert r.status_code == 200, r.text
    jid = r.json()["id"]
    j = esperar(cliente, jid)
    assert j["estado"] == "listo" and len(j["versiones"]) == 1
    assert j["fotos"] == ["Fisuras/a.jpg", "b.jpg"]
    bl = j["versiones"][0]["spec"]["blocks"]
    # foto inexistente descartada, foto no usada va al anexo
    assert [b["file"] for b in bl if b["type"] == "photo"] == ["Fisuras/a.jpg", "b.jpg"]
    assert any(b.get("text") == "Anexo fotográfico" for b in bl)

    assert cliente.get(f"/api/trabajos/{jid}/fotos/Fisuras/a.jpg").status_code == 200
    assert cliente.get(f"/api/trabajos/{jid}/fotos/..%2Fjob.json").status_code == 404

    # modificar -> v2 con feedback
    assert cliente.post(f"/api/trabajos/{jid}/modificar", data={"cambios": "más corto"}).status_code == 200
    j = esperar(cliente, jid)
    assert len(j["versiones"]) == 2 and cliente.llamadas[-1]["feedback"] == "más corto"
    assert cliente.llamadas[-1]["previo"] is not None

    d = cliente.get(f"/api/trabajos/{jid}/descargar")
    assert d.status_code == 200
    doc = Document(io.BytesIO(d.content))
    assert any("Informe test" in p.text for p in doc.paragraphs)
    assert len(doc.inline_shapes) == 2 and len(doc.tables) == 1

    # aprobar -> plantilla
    assert cliente.get("/api/plantillas").json() == []
    assert cliente.post(f"/api/trabajos/{jid}/aprobar").status_code == 200
    pl = cliente.get("/api/plantillas").json()
    assert len(pl) == 1 and pl[0]["nombre"] == "Puente"
    assert cliente.post(f"/api/trabajos/{jid}/modificar", data={"cambios": "x"}).status_code == 409

    # utilizar existente: el informe aprobado llega como ejemplo
    r = crear(cliente, plantilla_id=pl[0]["id"], nombre="Puente 2")
    assert r.status_code == 200, r.text
    esperar(cliente, r.json()["id"])
    assert cliente.llamadas[-1]["ejemplo"]["title"] == "Informe test"
    assert cliente.delete(f"/api/plantillas/{pl[0]['id']}").status_code == 200


def test_validaciones(cliente):
    assert cliente.post("/api/trabajos", data={"nombre": "x", "instrucciones": "y"}).status_code == 400
    assert cliente.post("/api/trabajos", data={"nombre": "x", "instrucciones": "y", "plantilla_id": "0" * 12}
                        ).status_code == 404
    r = cliente.post("/api/trabajos", data={"nombre": "x", "instrucciones": "y"},
                     files=[("plantilla_archivo", ("p.docx", _plantilla(), "x"))])
    assert r.status_code == 400  # sin fotos
